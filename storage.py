"""Camada de armazenamento dos leads e respostas do diagnostico de marketing.

Estrategia:
- Se GOOGLE_SHEET_ID estiver definido e houver credenciais do Google
  (GOOGLE_CREDENTIALS_JSON ou credentials/google_service_account.json) -> Google Sheets.
- Caso contrario -> grava localmente em data/leads.xlsx (funciona offline, sem setup).

Em ambos os casos os dados podem ser exportados em .xlsx pelo painel /admin.
"""
import os
import threading
import datetime as dt

from openpyxl import Workbook, load_workbook

from config import CADASTRO_FIELDS, iter_questions

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
CRED_PATH = os.path.join(BASE_DIR, "credentials", "google_service_account.json")
LOCAL_PATH = os.path.join(DATA_DIR, "leads.xlsx")
SHEET_ENV = "GOOGLE_SHEET_ID"

os.makedirs(DATA_DIR, exist_ok=True)

_lock = threading.Lock()

# Cabecalhos das duas "abas" (uma para o cadastro/lead, outra para as respostas).
LEAD_HEADERS = ["id", "data_hora"] + [f["name"] for f in CADASTRO_FIELDS]
RESP_HEADERS = ["id", "data_hora", "grupo", "pergunta", "resposta"]


# ---------------------------------------------------------------------------
# Google Sheets (opcional)
# ---------------------------------------------------------------------------
def _load_google_creds(scopes):
    """Carrega as credenciais da conta de servico.

    Em producao (Render), o JSON vem da variavel de ambiente
    GOOGLE_CREDENTIALS_JSON. Localmente, do arquivo em credentials/.
    """
    from google.oauth2.service_account import Credentials

    raw = os.environ.get("GOOGLE_CREDENTIALS_JSON")
    if raw:
        import json
        return Credentials.from_service_account_info(json.loads(raw), scopes=scopes)
    if os.path.exists(CRED_PATH):
        return Credentials.from_service_account_file(CRED_PATH, scopes=scopes)
    return None


def _google_client():
    """Retorna (worksheet_leads, worksheet_respostas) ou None se nao configurado."""
    sheet_id = os.environ.get(SHEET_ENV)
    if not sheet_id:
        return None
    try:
        import gspread

        scopes = ["https://www.googleapis.com/auth/spreadsheets"]
        creds = _load_google_creds(scopes)
        if creds is None:
            return None
        gc = gspread.authorize(creds)
        sh = gc.open_by_key(sheet_id)

        def _tab(name, headers):
            try:
                ws = sh.worksheet(name)
            except gspread.WorksheetNotFound:
                ws = sh.add_worksheet(title=name, rows=1000, cols=len(headers) + 5)
                ws.append_row(headers)
            if not ws.row_values(1):
                ws.append_row(headers)
            return ws

        return _tab("Leads", LEAD_HEADERS), _tab("Respostas", RESP_HEADERS)
    except Exception as e:  # pragma: no cover - depende de credenciais externas
        print(f"[storage] Google Sheets indisponivel, usando arquivo local: {e}")
        return None


def storage_backend():
    return "Google Sheets" if _google_client() else "Arquivo Excel local"


# ---------------------------------------------------------------------------
# Excel local
# ---------------------------------------------------------------------------
def _ensure_local_wb():
    if os.path.exists(LOCAL_PATH):
        return load_workbook(LOCAL_PATH)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Leads"
    ws1.append(LEAD_HEADERS)
    ws2 = wb.create_sheet("Respostas")
    ws2.append(RESP_HEADERS)
    wb.save(LOCAL_PATH)
    return wb


# ---------------------------------------------------------------------------
# API publica
# ---------------------------------------------------------------------------
def new_id():
    return dt.datetime.now().strftime("%Y%m%d%H%M%S%f")


def save_lead(lead_id, data):
    """Grava o cadastro (contato)."""
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row = [lead_id, now] + [str(data.get(f["name"], "")) for f in CADASTRO_FIELDS]
    with _lock:
        g = _google_client()
        if g:
            g[0].append_row(row, value_input_option="RAW")
        else:
            wb = _ensure_local_wb()
            wb["Leads"].append(row)
            wb.save(LOCAL_PATH)


def save_respostas(lead_id, answers):
    """Grava as respostas do questionario. answers: {qid: texto}."""
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows = []
    for qid, grupo_id, grupo_titulo, texto in iter_questions():
        rows.append([lead_id, now, grupo_titulo, texto, str(answers.get(qid, "")).strip()])
    with _lock:
        g = _google_client()
        if g:
            g[1].append_rows(rows, value_input_option="RAW")
        else:
            wb = _ensure_local_wb()
            ws = wb["Respostas"]
            for r in rows:
                ws.append(r)
            wb.save(LOCAL_PATH)


def export_xlsx_bytes():
    """Retorna um .xlsx (bytes) com tudo que esta armazenado, para download no /admin."""
    import io

    g = _google_client()
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Leads"
    ws2 = wb.create_sheet("Respostas")

    if g:
        leads = g[0].get_all_values()
        resp = g[1].get_all_values()
        for r in (leads or [LEAD_HEADERS]):
            ws1.append(r)
        for r in (resp or [RESP_HEADERS]):
            ws2.append(r)
    else:
        src = _ensure_local_wb()
        ws1.delete_rows(1, ws1.max_row)
        for r in src["Leads"].iter_rows(values_only=True):
            ws1.append(list(r))
        for r in src["Respostas"].iter_rows(values_only=True):
            ws2.append(list(r))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
